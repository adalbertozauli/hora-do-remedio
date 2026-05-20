from pathlib import Path

from docx import Document


def extract_text_from_docx(file_path: str) -> str:
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError("Arquivo não encontrado.")
    if path.suffix.lower() != ".docx":
        raise ValueError("Selecione um arquivo no formato .docx.")

    document = Document(str(path))
    lines: list[str] = []

    for paragraph in document.paragraphs:
        text = " ".join(paragraph.text.split())
        if text:
            lines.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [" ".join(cell.text.split()) for cell in row.cells]
            line = " | ".join(cell for cell in cells if cell)
            if line:
                lines.append(line)

    extracted = "\n".join(lines).strip()
    if not extracted:
        raise ValueError("Não foi possível encontrar texto legível no documento.")
    return extracted
